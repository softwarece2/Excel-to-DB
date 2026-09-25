"""
excel/parser.py
----------------
Parser for the current ATE specification workbook layout.

Expected columns:
Sl. No. | Param code | Main type | Parameter type | Parameter description |
TY clause no | Minimum value | maximum value | Unit | Test Type |
Test by QC & CQAI | Duration in Minutes
"""

from dataclasses import dataclass
from typing import Dict, List, Optional

from openpyxl.worksheet.worksheet import Worksheet


def _normalize(text) -> str:
    """Collapse whitespace/newlines and lowercase for robust header matching."""
    if text is None:
        return ""
    return " ".join(str(text).split()).strip().lower()


@dataclass
class ParameterRow:
    sl_no: Optional[int]
    param_code: Optional[str]
    main_type: Optional[str]
    param_type: Optional[str]
    spec_val: Optional[str]
    ty_clause_no: Optional[str]
    min_raw: object = None
    max_raw: object = None
    unit_raw: object = None
    test_type: Optional[str] = None
    test_by: Optional[str] = None
    duration_raw: object = None


def _locate_header_columns(
    ws: Worksheet, header_row: int, columns_map: Dict[str, str]
) -> Dict[str, int]:
    normalized_targets = {
        key: _normalize(label) for key, label in columns_map.items()
    }

    found: Dict[str, int] = {}
    for col_idx in range(1, ws.max_column + 1):
        normalized_cell = _normalize(ws.cell(row=header_row, column=col_idx).value)
        if not normalized_cell:
            continue

        for key, target in normalized_targets.items():
            if key not in found and normalized_cell == target:
                found[key] = col_idx

    missing = set(normalized_targets) - set(found)
    if missing:
        raise ValueError(
            f"Sheet '{ws.title}': could not locate header column(s) {sorted(missing)} "
            f"in row {header_row}. Check mappings/column_mapping.json."
        )
    return found


def _to_int(value) -> Optional[int]:
    if value in (None, ""):
        return None
    if isinstance(value, bool):
        raise ValueError(f"Boolean value '{value}' is not valid for Sl. No.")
    try:
        return int(float(str(value).strip()))
    except (TypeError, ValueError):
        raise ValueError(f"Could not convert Sl. No. value '{value}' to integer.")


def _to_text(value) -> Optional[str]:
    if value in (None, ""):
        return None
    return str(value).strip()


def parse_parameter_rows(
    ws: Worksheet,
    header_row: int,
    data_start_row: int,
    columns_map: Dict[str, str],
) -> List[ParameterRow]:
    """
    Read all populated parameter rows.

    Main type / Parameter type may be merged in Excel, so the parser
    forward-fills those values down just like Sl. No. in the old template.
    """
    col_idx = _locate_header_columns(ws, header_row, columns_map)

    rows: List[ParameterRow] = []
    last_main_type = None
    last_param_type = None

    for row_idx in range(data_start_row, ws.max_row + 1):
        raw = {
            key: ws.cell(row=row_idx, column=idx).value
            for key, idx in col_idx.items()
        }

        if raw["main_type"] not in (None, ""):
            last_main_type = _to_text(raw["main_type"])
        if raw["param_type"] not in (None, ""):
            last_param_type = _to_text(raw["param_type"])

        # Parameter description is the required content field.
        if raw["spec_val"] in (None, ""):
            continue

        rows.append(
            ParameterRow(
                sl_no=_to_int(raw["sl_no"]),
                param_code=_to_text(raw["param_code"]),
                main_type=last_main_type,
                param_type=last_param_type,
                spec_val=_to_text(raw["spec_val"]),
                ty_clause_no=_to_text(raw["ty_clause_no"]),
                min_raw=raw["min"],
                max_raw=raw["max"],
                unit_raw=raw["unit"],
                test_type=_to_text(raw["test_type"]),
                test_by=_to_text(raw["test_by"]),
                duration_raw=raw["duration"],
            )
        )

    return rows
